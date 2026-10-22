"""Excel and CSV loading, header detection, raw layer, input errors."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from conftest import HEADERS, make_row, write_csv, write_xlsx
from scholarpath_data.errors import SourceError
from scholarpath_data.load_excel import (
    check_input_path,
    discover_input_files,
    load_source,
    resolve_input,
    sha256_file,
    write_raw_snapshot,
)


def test_loads_excel_rows_columns_and_source_row_numbers(sample_xlsx: Path) -> None:
    source = load_source(sample_xlsx)
    assert source.file_format == "xlsx"
    assert [s.name for s in source.sheets] == ["Sheet1"]
    sheet = source.sheets[0]
    assert sheet.columns == HEADERS
    assert len(sheet.frame) == 3
    assert list(sheet.frame.index) == [2, 3, 4]  # Excel row numbers (header is row 1)
    assert sheet.frame.iloc[0]["Scholarship Name"] == "Alpha Merit Scholarship"


def test_empty_sheet_is_skipped_and_reported(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "multi.xlsx", [make_row()], extra_sheets={"central": []})
    source = load_source(path)
    assert [s.name for s in source.sheets] == ["Sheet1"]
    assert source.empty_sheets == ["central"]


def test_trailing_empty_columns_and_blank_rows_are_dropped(tmp_path: Path) -> None:
    path = tmp_path / "wide.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "State"])
    ws.append(["A", "MP"])
    ws.append([None, None])  # blank row
    ws.append(["B", "UP"])
    ws.cell(row=5, column=60, value=None)  # extends the declared dimensions with empty cells
    wb.save(path)
    sheet = load_source(path).sheets[0]
    assert sheet.columns == ["Name", "State"]
    assert list(sheet.frame.index) == [2, 4]  # blank row 3 skipped, numbering preserved
    assert sheet.blank_rows_skipped == 2  # rows 3 and 5 (the cell at column 60 made row 5 exist but empty)


def test_header_is_detected_below_a_title_row(tmp_path: Path) -> None:
    path = tmp_path / "titled.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Scholarship master list 2026"])
    ws.append([])
    ws.append(["Name", "State"])
    ws.append(["A", "MP"])
    wb.save(path)
    sheet = load_source(path).sheets[0]
    assert sheet.header_row == 3
    assert list(sheet.frame.index) == [4]


def test_repeated_headers_are_made_unique(tmp_path: Path) -> None:
    path = tmp_path / "dup.csv"
    path.write_text("Name,name,State\nA,B,MP\n", encoding="utf-8")
    source = load_source(path)
    assert source.sheets[0].columns == ["Name", "name__2", "State"]
    assert any("repeated column header" in n for n in source.notes)


def test_loads_csv_utf8_with_bom_and_ruppee_sign(tmp_path: Path) -> None:
    path = write_csv(tmp_path / "s.csv", [make_row(**{"key benefits": "₹5,000 per month"})])
    source = load_source(path)
    assert source.file_format == "csv"
    assert source.sheets[0].frame.iloc[0]["key benefits"] == "₹5,000 per month"
    assert source.sheets[0].columns[0] == "SCHOLARSHIP_ID"  # BOM must not leak into the header


def test_csv_windows_1252_fallback_and_semicolon_delimiter(tmp_path: Path) -> None:
    path = tmp_path / "w.csv"
    path.write_bytes("Name;Notes\nCaf\xe9 Fund;ok\n".encode("cp1252"))
    source = load_source(path)
    assert source.sheets[0].frame.iloc[0]["Name"] == "Café Fund"
    assert any("cp1252" in n for n in source.notes)


def test_csv_ragged_rows_do_not_crash(tmp_path: Path) -> None:
    path = tmp_path / "r.csv"
    path.write_text("A,B,C\n1,2\n3,4,5,6\n", encoding="utf-8")
    sheet = load_source(path).sheets[0]
    assert len(sheet.frame) == 2


def test_sheet_selection_and_unknown_sheet_error(sample_xlsx: Path) -> None:
    assert [s.name for s in load_source(sample_xlsx, ["Sheet1"]).sheets] == ["Sheet1"]
    with pytest.raises(SourceError, match="Sheet\\(s\\) not found.*Nope.*Available: Sheet1"):
        load_source(sample_xlsx, ["Nope"])


def test_missing_file_gives_actionable_message(tmp_path: Path) -> None:
    with pytest.raises(SourceError, match="(?s)not found.*quotes"):
        check_input_path(tmp_path / "nope.xlsx")


def test_folder_and_unsupported_extension_are_rejected(tmp_path: Path) -> None:
    with pytest.raises(SourceError, match="folder"):
        check_input_path(tmp_path)
    legacy = tmp_path / "old.xls"
    legacy.write_bytes(b"x")
    with pytest.raises(SourceError, match="Unsupported.*Save As > .xlsx"):
        check_input_path(legacy)
    other = tmp_path / "data.json"
    other.write_text("{}")
    with pytest.raises(SourceError, match="Unsupported"):
        check_input_path(other)


def test_corrupt_xlsx_is_reported_not_crashed(tmp_path: Path) -> None:
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is not a zip file")
    with pytest.raises(SourceError, match="Cannot read bad.xlsx as an Excel workbook"):
        load_source(bad)


def test_completely_empty_workbook_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(SourceError, match="contains no data"):
        load_source(path)


def test_loading_never_modifies_the_source_file(sample_xlsx: Path) -> None:
    before_hash, before_mtime = sha256_file(sample_xlsx), sample_xlsx.stat().st_mtime_ns
    load_source(sample_xlsx)
    assert sha256_file(sample_xlsx) == before_hash
    assert sample_xlsx.stat().st_mtime_ns == before_mtime


def test_discovery_finds_single_candidate_and_ignores_temp_files(tmp_path: Path, sample_xlsx: Path) -> None:
    folder = sample_xlsx.parent
    (folder / "~$sample.xlsx").write_bytes(b"lock")
    (folder / "notes.txt").write_text("x")
    assert discover_input_files([folder]) == [sample_xlsx.resolve()]
    assert resolve_input(None, [folder]) == sample_xlsx.resolve()


def test_discovery_errors_when_none_or_ambiguous(tmp_path: Path, sample_xlsx: Path) -> None:
    with pytest.raises(SourceError, match="No --input given"):
        resolve_input(None, [tmp_path / "empty-dir"])
    write_csv(sample_xlsx.parent / "other.csv", [make_row()])
    with pytest.raises(SourceError, match="Several candidate files"):
        resolve_input(None, [sample_xlsx.parent])


def test_raw_snapshot_is_verbatim_and_extracted_at_is_stable(tmp_path: Path) -> None:
    from datetime import datetime, timedelta, timezone

    path = write_xlsx(tmp_path / "raw_src.xlsx", [make_row(max_family_income="<4,50,000", min_percentage=80)])
    source = load_source(path)
    t1 = datetime(2026, 1, 1, tzinfo=timezone.utc)
    first = write_raw_snapshot(source, tmp_path / "raw", t1)
    again = write_raw_snapshot(source, tmp_path / "raw", t1 + timedelta(days=5))
    assert first.extracted_at == again.extracted_at == t1  # same content -> same first-extraction time
    assert first.directory == again.directory
    assert first.directory.name == f"raw_src-{source.sha256[:8]}"

    text = first.csv_files["Sheet1"].read_text(encoding="utf-8-sig")
    assert "<4,50,000" in text and ",80," in text  # original text, nothing cleaned
    manifest = json.loads(first.manifest_path.read_text(encoding="utf-8"))
    assert manifest["sha256"] == source.sha256
    assert manifest["sheets"][0]["data_rows"] == 1


def test_changed_source_gets_a_new_raw_folder_instead_of_overwriting(tmp_path: Path) -> None:
    path = write_xlsx(tmp_path / "v.xlsx", [make_row()])
    first = write_raw_snapshot(load_source(path), tmp_path / "raw")
    write_xlsx(path, [make_row(), make_row(**{"Scholarship Name": "Second"})])
    second = write_raw_snapshot(load_source(path), tmp_path / "raw")
    assert first.directory != second.directory
    assert first.directory.is_dir() and second.directory.is_dir()
